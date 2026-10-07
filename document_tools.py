import sys, json, os, csv
import base64
from io import BytesIO

def extract(ext, data):
    if ext in ('.txt','.md','.log'):
        return data.decode('utf-8-sig',errors='replace')
    if ext == '.csv':
        rows=list(csv.reader(data.decode('utf-8-sig',errors='replace').splitlines()))
        rows=[row for row in rows if any(cell.strip() for cell in row)]
        if not rows: return ''
        width=max(map(len,rows)); rows=[row+['']*(width-len(row)) for row in rows]
        cells=lambda row:' | '.join(str(cell).strip().replace('|','\\|').replace('\n','<br>') for cell in row)
        return '| '+cells(rows[0])+' |\n| '+' | '.join('---' for _ in range(width))+' |\n'+'\n'.join('| '+cells(row)+' |' for row in rows[1:])
    if ext == '.json':
        value=json.loads(data.decode('utf-8-sig',errors='replace'))
        return '```json\n'+json.dumps(value,ensure_ascii=False,indent=2)+'\n```'
    if ext == '.html':
        from html.parser import HTMLParser
        class TextExtractor(HTMLParser):
            def __init__(self): super().__init__(); self.parts=[]; self.hidden=0
            def handle_starttag(self,tag,attrs):
                if tag in ('script','style','noscript'): self.hidden+=1
                elif not self.hidden and tag=='li': self.parts.extend(['\n','- '])
                elif not self.hidden and tag in ('h1','h2','h3','h4','p','div','tr','br','section','article'): self.parts.append('\n')
            def handle_endtag(self,tag):
                if tag in ('script','style','noscript') and self.hidden: self.hidden-=1
                elif not self.hidden and tag in ('h1','h2','h3','h4','p','div','li','tr','section','article'): self.parts.append('\n')
            def handle_data(self,text):
                if not self.hidden: self.parts.append(text)
        parser=TextExtractor(); parser.feed(data.decode('utf-8',errors='replace'))
        return '\n'.join(line.strip() for line in ''.join(parser.parts).splitlines() if line.strip())
    if ext == '.pdf':
        from pypdf import PdfReader
        from io import BytesIO
        pages=[]
        for index, page in enumerate(PdfReader(BytesIO(data)).pages, 1):
            try: extracted=page.extract_text(extraction_mode='layout')
            except TypeError: extracted=page.extract_text()
            text=(extracted or '').strip()
            pages.append(f'## Page {index}\n\n{text}' if text else f'## Page {index}\n\n[No selectable text found on this page.]')
        return '\n\n'.join(pages)
    if ext == '.docx':
        from docx import Document
        from io import BytesIO
        doc=Document(BytesIO(data))
        blocks=[]
        for child in doc.element.body:
            tag=child.tag.rsplit('}',1)[-1]
            if tag == 'p':
                from docx.text.paragraph import Paragraph
                p=Paragraph(child,doc)
                text=p.text.strip()
                if not text: continue
                style=(p.style.name if p.style else '').lower()
                if style.startswith('heading'):
                    level=''.join(c for c in style if c.isdigit()) or '2'
                    blocks.append('#'*min(int(level),6)+' '+text)
                elif 'title' in style: blocks.append('# '+text)
                elif style.startswith('list') or 'bullet' in style: blocks.append('- '+text)
                else: blocks.append(text)
            elif tag == 'tbl':
                from docx.table import Table
                table=Table(child,doc)
                rows=[[' '.join(cell.text.split()) for cell in row.cells] for row in table.rows]
                while rows and not any(rows[0]): rows.pop(0)
                if not rows: continue
                width=max(map(len,rows)); rows=[row+['']*(width-len(row)) for row in rows]
                header=rows[0]
                blocks.append('| '+' | '.join(v.replace('|','\\|') for v in header)+' |')
                blocks.append('| '+' | '.join('---' for _ in header)+' |')
                blocks.extend('| '+' | '.join(v.replace('|','\\|') for v in row)+' |' for row in rows[1:])
        return '\n\n'.join(blocks)
    if ext == '.xlsx':
        from openpyxl import load_workbook
        from io import BytesIO
        wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
        blocks=[]
        for ws in wb.worksheets:
            rows=[]
            for row in ws.iter_rows(values_only=True):
                values=[str(v).strip() if v is not None else '' for v in row]
                while values and not values[-1]: values.pop()
                if any(values): rows.append(values)
            if not rows: continue
            width=max(map(len,rows)); rows=[row+['']*(width-len(row)) for row in rows]
            blocks.append(f'# Sheet: {ws.title} · {len(rows)} rows')
            blocks.append('| '+' | '.join(v.replace('|','\\|') for v in rows[0])+' |')
            blocks.append('| '+' | '.join('---' for _ in range(width))+' |')
            blocks.extend('| '+' | '.join(v.replace('|','\\|') for v in row)+' |' for row in rows[1:])
        wb.close()
        return '\n\n'.join(blocks)
    raise ValueError('Unsupported file format')

def render_pdf_pages(data, limit=4):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    page_count = len(pdf)
    pages = []
    for index in range(min(page_count, limit)):
        page = pdf[index]
        width, height = page.get_size()
        scale = min(1.25, 1400 / max(width, height, 1))
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil().convert('RGB')
        image.thumbnail((1400, 1400))
        buffer = BytesIO()
        image.save(buffer, format='JPEG', quality=78, optimize=True)
        pages.append({'page': index + 1, 'image': base64.b64encode(buffer.getvalue()).decode('ascii')})
    pdf.close()
    return {'page_count': page_count, 'pages': pages}

def write_docx(out, data):
    from docx import Document
    x=json.loads(data.decode('utf-8'))
    doc=Document()
    doc.add_heading(x.get('title') or 'Approval Note',0)
    doc.add_paragraph('CONFIDENTIAL · LOCAL WORKBENCH DRAFT')
    for block in x.get('body','').split('\n'):
        line=block.strip()
        if not line: continue
        if line.startswith('# '): doc.add_heading(line[2:],1)
        elif line.startswith('## '): doc.add_heading(line[3:],2)
        elif line.startswith('- ') or line.startswith('* '): doc.add_paragraph(line[2:],style='List Bullet')
        else: doc.add_paragraph(line)
    doc.save(out)

if __name__=='__main__':
    mode=sys.argv[1]
    if mode=='extract': sys.stdout.write(extract(sys.argv[2],sys.stdin.buffer.read()))
    elif mode=='render-pdf-pages': sys.stdout.write(json.dumps(render_pdf_pages(sys.stdin.buffer.read())))
    elif mode=='write-docx': write_docx(sys.argv[2],sys.stdin.buffer.read())
    else: raise ValueError('Unknown helper command')
